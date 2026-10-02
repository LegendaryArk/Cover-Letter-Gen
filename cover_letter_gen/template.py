"""Find and fill {{placeholders}} in a .docx while preserving formatting.

Placeholder syntax:
  {{company}}            simple field, value extracted from the job description
  {{? guidance text}}    written slot, Claude writes 1-2 sentences following the guidance

Google Docs exports often split one placeholder across several runs, so matching
is done on the paragraph's concatenated text and the replacement is written into
the run where the marker starts (inheriting that run's font, size, bold, etc.).
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, asdict
from typing import Iterator

from docx import Document
from docx.document import Document as DocumentT
from docx.table import Table
from docx.text.paragraph import Paragraph

MARKER_RE = re.compile(r"\{\{(.+?)\}\}", re.DOTALL)


@dataclass
class Slot:
    id: str          # field name, or "w1", "w2", ... for written slots
    kind: str        # "field" | "written"
    key: str         # normalized marker contents, used to match on fill
    guidance: str = ""

    def to_dict(self) -> dict:
        return asdict(self)


def normalize_key(inner: str) -> str:
    inner = " ".join(inner.split())
    if inner.startswith("?"):
        return "?" + inner[1:].strip()
    return inner.lower().replace(" ", "_")


def _iter_table_paragraphs(table: Table) -> Iterator[Paragraph]:
    for row in table.rows:
        for cell in row.cells:
            yield from cell.paragraphs
            for nested in cell.tables:
                yield from _iter_table_paragraphs(nested)


def _iter_container(container) -> Iterator[Paragraph]:
    yield from container.paragraphs
    for table in container.tables:
        yield from _iter_table_paragraphs(table)


def iter_paragraphs(doc: DocumentT, include_headers: bool = True) -> Iterator[Paragraph]:
    yield from _iter_container(doc)
    if not include_headers:
        return
    seen = set()
    for section in doc.sections:
        for part in (
            section.header, section.first_page_header, section.even_page_header,
            section.footer, section.first_page_footer, section.even_page_footer,
        ):
            if part.is_linked_to_previous or id(part._element) in seen:
                continue
            seen.add(id(part._element))
            yield from _iter_container(part)


def load(data: bytes) -> DocumentT:
    return Document(io.BytesIO(data))


def extract_slots(doc: DocumentT) -> list[Slot]:
    slots: list[Slot] = []
    seen: set[str] = set()
    n_written = 0
    for p in iter_paragraphs(doc):
        for m in MARKER_RE.finditer(p.text):
            key = normalize_key(m.group(1))
            if not key or key == "?" or key in seen:
                continue
            seen.add(key)
            if key.startswith("?"):
                n_written += 1
                slots.append(Slot(id=f"w{n_written}", kind="written", key=key, guidance=key[1:].strip()))
            else:
                slots.append(Slot(id=key, kind="field", key=key))
    return slots


def template_text(doc: DocumentT) -> str:
    """Body text of the template, with written slots shown as [WRITE w1: ...]."""
    slots = {s.key: s for s in extract_slots(doc)}

    def show(m: re.Match) -> str:
        s = slots.get(normalize_key(m.group(1)))
        if s is None:
            return m.group(0)
        return f"[WRITE {s.id}: {s.guidance}]" if s.kind == "written" else f"{{{{{s.id}}}}}"

    return "\n".join(MARKER_RE.sub(show, p.text) for p in iter_paragraphs(doc, include_headers=False))


def _fill_paragraph(p: Paragraph, values: dict[str, str]) -> None:
    runs = p.runs
    if not runs:
        return
    texts = [r.text for r in runs]
    full = "".join(texts)
    if "{{" not in full:
        return
    # Map each character offset to (run index, offset within run).
    spans = []  # (start, end) of each run in the original concatenated text
    pos = 0
    for t in texts:
        spans.append((pos, pos + len(t)))
        pos += len(t)

    def locate(offset: int, end: bool = False) -> tuple[int, int]:
        # Start offsets belong to the run containing that character; end offsets
        # (exclusive) to the run containing the character just before them.
        for i, (a, b) in enumerate(spans):
            if (a <= offset < b) if not end else (a < offset <= b):
                return i, offset - a
        raise ValueError(offset)

    matches = [m for m in MARKER_RE.finditer(full) if normalize_key(m.group(1)) in values]
    # Replace back to front so earlier offsets stay valid.
    for m in reversed(matches):
        replacement = values[normalize_key(m.group(1))]
        si, so = locate(m.start())
        ei, eo = locate(m.end(), end=True)
        if si == ei:
            t = texts[si]
            texts[si] = t[:so] + replacement + t[eo:]
        else:
            texts[si] = texts[si][:so] + replacement
            for k in range(si + 1, ei):
                texts[k] = ""
            texts[ei] = texts[ei][eo:]
    for r, t in zip(runs, texts):
        if r.text != t:
            r.text = t


def fill(doc: DocumentT, values: dict[str, str]) -> DocumentT:
    """Replace markers whose normalized key is in `values`. Mutates and returns doc."""
    for p in iter_paragraphs(doc):
        _fill_paragraph(p, values)
    return doc


def fill_bytes(template: bytes, values: dict[str, str]) -> bytes:
    doc = fill(load(template), values)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
