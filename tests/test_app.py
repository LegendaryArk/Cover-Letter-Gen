import io
import json

import pytest
from docx import Document
from fastapi.testclient import TestClient

from cover_letter_gen import app as app_module
from cover_letter_gen import claude_client as cc
from cover_letter_gen import export, storage, template


def make_template() -> bytes:
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "{{my_name}} · {{date}}"
    p = doc.add_paragraph()
    p.add_run("Dear {{hiring")           # marker split across differently formatted runs
    bold = p.add_run("_manager}},")
    bold.bold = True
    p2 = doc.add_paragraph("I am applying for the {{role}} role at {{company}}. ")
    p2.add_run("{{? 1 sentence on why their product excites me. Pick one: mission / tech}}").italic = True
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Location: {{ Location }}"
    doc.add_paragraph("{{company}} again. {{? one sentence on team fit}}")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_extract_slots():
    slots = template.extract_slots(template.load(make_template()))
    by_id = {s.id: s for s in slots}
    assert set(by_id) == {"my_name", "date", "hiring_manager", "role", "company", "location", "w1", "w2"}
    assert by_id["w1"].kind == "written"
    assert by_id["w1"].guidance.startswith("1 sentence on why")
    assert by_id["w2"].guidance == "one sentence on team fit"


def test_template_text_marks_written_slots():
    text = template.template_text(template.load(make_template()))
    assert "[WRITE w1: 1 sentence on why" in text
    assert "{{company}}" in text


def test_fill_preserves_formatting_and_covers_everywhere():
    doc = template.load(make_template())
    slots = template.extract_slots(doc)
    values = {s.key: f"<{s.id}>" for s in slots}
    out = template.load(template.fill_bytes(make_template(), values))
    all_text = "\n".join(p.text for p in template.iter_paragraphs(out))
    assert "{{" not in all_text
    paras = out.paragraphs
    assert paras[0].text == "Dear <hiring_manager>,"
    # Replacement lives in the run where the marker started (not bold); trailing comma stays bold.
    assert paras[0].runs[0].text == "Dear <hiring_manager>"
    assert paras[0].runs[0].bold is None
    assert paras[0].runs[1].text == "," and paras[0].runs[1].bold
    assert paras[1].runs[-1].text == "<w1>" and paras[1].runs[-1].italic
    assert paras[2].text == "<company> again. <w2>"
    assert out.tables[0].cell(0, 0).text == "Location: <location>"
    assert out.sections[0].header.paragraphs[0].text == "<my_name> · <date>"


def test_build_filename():
    s = storage.Settings(first_name="Noah", last_name="Sun")
    assert export.build_filename(s, "Stripe, Inc.", "Software Engineer Intern (Summer 2027)") == \
        "Noah_Sun_CoverLetter_Stripe_Inc_Software_Engineer_Intern_Summer_2027.pdf"
    assert export.build_filename(storage.Settings(), "Acme", "Dev") == "CoverLetter_Acme_Dev.pdf"


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(storage, "DATA_DIR", tmp_path / "data")
    monkeypatch.setattr(storage, "OUTPUT_DIR", tmp_path / "output")
    monkeypatch.setattr(storage, "TEMPLATE_PATH", tmp_path / "data" / "template.docx")
    monkeypatch.setattr(storage, "RESUME_PATH", tmp_path / "data" / "resume.txt")
    monkeypatch.setattr(storage, "RESUME_PDF_PATH", tmp_path / "data" / "resume.pdf")
    monkeypatch.setattr(storage, "SETTINGS_PATH", tmp_path / "data" / "settings.json")
    monkeypatch.setattr(export, "DATA_DIR", tmp_path / "data")
    storage.ensure_dirs()
    app_module._research_cache.clear()

    calls = {"research": 0, "write": []}

    async def fake_extract(jd, field_names):
        role = jd.split("\n")[0]
        return cc.JobInfo(company="Stripe, Inc.", role=role, company_short="Stripe", role_short=role,
                          fields=[cc.FieldValue(name=n, value=f"{n}-val") for n in field_names])

    async def fake_research(company, roles, jd):
        calls["research"] += 1
        return f"{company} builds payments. Roles: {', '.join(roles)}"

    async def fake_write(**kw):
        calls["write"].append(kw)
        return {s["id"]: f"Text for {s['id']}" for s in kw["slots"]}

    monkeypatch.setattr(cc, "extract_job", fake_extract)
    monkeypatch.setattr(cc, "research_company", fake_research)
    monkeypatch.setattr(cc, "write_slots", fake_write)
    c = TestClient(app_module.app)
    c.calls = calls
    return c


def test_generate_and_export_flow(client):
    r = client.post("/api/template", files={"file": ("t.docx", make_template())})
    assert r.status_code == 200
    client.put("/api/settings", json={**storage.Settings().model_dump(), "first_name": "Noah", "last_name": "Sun"})
    client.post("/api/resume", files={"file": ("resume.txt", b"Built things at places.")})

    r = client.post("/api/generate", json={"jds": ["Backend Engineer\n...", "Frontend Engineer\n..."],
                                           "include_resume": False})
    events = [json.loads(line) for line in r.text.splitlines()]
    jobs = [e["job"] for e in events if e["type"] == "job"]
    assert events[-1]["type"] == "done"
    assert len(jobs) == 2 and not any(j["error"] for j in jobs)
    assert client.calls["research"] == 1  # shared across same-company roles
    assert all(kw["resume"] is None for kw in client.calls["write"])
    assert {tuple(j["siblings"]) for j in jobs} == {("Frontend Engineer",), ("Backend Engineer",)}
    assert jobs[0]["fields"]["hiring_manager"] == "hiring_manager-val"
    assert "my_name" not in jobs[0]["fields"]  # filled from settings, not Claude

    if not _has_soffice():
        pytest.skip("LibreOffice not installed")
    r = client.post("/api/export", json={"jobs": jobs})
    assert r.status_code == 200, r.text
    names = sorted(f["filename"] for f in r.json()["files"])
    assert names == ["Noah_Sun_CoverLetter_Stripe_Backend_Engineer.pdf",
                     "Noah_Sun_CoverLetter_Stripe_Frontend_Engineer.pdf"]
    pdf = client.get(f"/api/files/{names[0]}").content
    assert pdf.startswith(b"%PDF")
    from pypdf import PdfReader
    text = PdfReader(io.BytesIO(pdf)).pages[0].extract_text()
    assert "Text for w1" in text and "Noah Sun" in text and "{{" not in text
    z = client.post("/api/zip", json={"filenames": names})
    assert z.headers["content-type"] == "application/zip"


def test_regenerate_slot_passes_hint_and_previous(client):
    client.post("/api/template", files={"file": ("t.docx", make_template())})
    job = {"jd": "Backend Engineer", "written": {"w1": "old"}, "research": "r"}
    r = client.post("/api/regenerate-slot", json={"job": job, "slot_id": "w1", "hint": "shorter"})
    assert r.json() == {"slot_id": "w1", "text": "Text for w1"}
    kw = client.calls["write"][-1]
    assert kw["hint"] == "shorter" and kw["previous"] == {"w1": "old"}


def test_path_traversal_blocked(client):
    assert client.get("/api/files/..%2Fdata%2Fsettings.json").status_code == 404


def make_pdf(text: str, pages: int = 1) -> bytes:
    """Minimal text PDF via LibreOffice, so pypdf can read the text back."""
    doc = Document()
    for i in range(pages):
        doc.add_paragraph(f"{text} page {i + 1}")
        if i < pages - 1:
            doc.add_page_break()
    buf = io.BytesIO()
    doc.save(buf)
    return export.docx_to_pdf(buf.getvalue())


def test_export_with_resume_attached(client):
    if not _has_soffice():
        pytest.skip("LibreOffice not installed")
    from pypdf import PdfReader
    client.post("/api/template", files={"file": ("t.docx", make_template())})
    client.put("/api/settings", json={**storage.Settings().model_dump(), "first_name": "Noah", "last_name": "Sun"})
    job = {"jd": "x", "company_short": "Stripe", "role_short": "Backend Engineer", "written": {"w1": "Letter body"}}

    # Text-only resume: usable as context but can't be attached.
    r = client.post("/api/resume", files={"file": ("resume.txt", b"Resume text")})
    assert r.json()["has_resume_pdf"] is False
    assert client.post("/api/export", json={"jobs": [job], "attach_resume": True}).status_code == 400

    r = client.post("/api/resume", files={"file": ("resume.pdf", make_pdf("RESUME CONTENT", pages=2))})
    assert r.json()["has_resume_pdf"] is True
    assert client.get("/api/state").json()["has_resume_pdf"] is True

    r = client.post("/api/export", json={"jobs": [job], "attach_resume": True})
    assert r.status_code == 200, r.text
    f = r.json()["files"][0]
    assert f["filename"] == "Noah_Sun_CoverLetter_Resume_Stripe_Backend_Engineer.pdf"
    pages = PdfReader(io.BytesIO(client.get(f["url"]).content)).pages
    assert len(pages) == 3
    assert "Letter body" in pages[0].extract_text()
    assert "RESUME CONTENT page 1" in pages[1].extract_text()
    assert "RESUME CONTENT page 2" in pages[2].extract_text()

    # Preview honours the toggle too; plain export is unchanged.
    prev = client.post("/api/preview?attach_resume=true", json=job).content
    assert len(PdfReader(io.BytesIO(prev)).pages) == 3
    plain = client.post("/api/export", json={"jobs": [job]}).json()["files"][0]["filename"]
    assert plain == "Noah_Sun_CoverLetter_Stripe_Backend_Engineer.pdf"


def test_docx_resume_converted_and_text_resume_clears_pdf(client):
    if not _has_soffice():
        pytest.skip("LibreOffice not installed")
    doc = Document()
    doc.add_paragraph("Docx resume")
    buf = io.BytesIO()
    doc.save(buf)
    assert client.post("/api/resume", files={"file": ("cv.docx", buf.getvalue())}).json()["has_resume_pdf"]
    assert storage.RESUME_PDF_PATH.read_bytes().startswith(b"%PDF")
    client.post("/api/resume", files={"file": ("cv.txt", b"New text resume")})
    assert not storage.RESUME_PDF_PATH.exists()


def _has_soffice():
    try:
        export.soffice_binary()
        return True
    except RuntimeError:
        return False
